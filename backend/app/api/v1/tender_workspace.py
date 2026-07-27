"""Integrated tender-preparation workspace backed by PostgreSQL 17."""

from __future__ import annotations

import asyncio
import hashlib
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import get_current_user
from app.db.base import get_async_session
from app.models.intelligence import KnowledgeEntry

router = APIRouter(prefix="/tender-workspace", tags=["tender-workspace"])


class WorkspaceStateUpdate(BaseModel):
    tasks: list[dict[str, Any]] | None = None
    document_statuses: dict[str, str] | None = None
    form_statuses: dict[str, str] | None = None
    approvals: list[dict[str, Any]] | None = None
    comments: list[dict[str, Any]] | None = None
    owner: str | None = None
    boq_completion: int | None = Field(None, ge=0, le=100)
    contractor_id: str | None = None


class ProfitMarginRequest(BaseModel):
    direct_cost_bdt: float = Field(..., ge=0)
    overhead_pct: float = Field(5, ge=0, le=100)
    tax_vat_pct: float = Field(7.5, ge=0, le=100)
    contingency_pct: float = Field(2, ge=0, le=100)
    target_profit_pct: float = Field(10, ge=0, le=100)


class BOQCompletionRequest(BaseModel):
    sor_agency: str = Field("BWDB", pattern="^(BWDB|LGED|PWD)$")
    zone: str | None = Field(None, pattern="^[A-Da-d]$")


class EstimateApprovalRequest(BaseModel):
    amount_bdt: float = Field(..., gt=0)
    source_record_id: str | None = None
    package_no: str | None = None
    rationale: str = "Approved APP recovery candidate"


class BOQResolutionRequest(BaseModel):
    item_id: str
    sor_code: str
    agency: str = Field(..., pattern="^(BWDB|LGED|PWD)$")
    zone: str = Field("A", pattern="^[A-Da-d]$")


class StageApprovalRequest(BaseModel):
    stage: str = Field(..., pattern="^(eligibility|documents|boq|pricing|compliance|final)$")
    status: str = Field("approved", pattern="^(approved|revoked)$")
    note: str | None = None


def _payload(data: Any) -> dict[str, Any]:
    if not isinstance(data, dict):
        return {}
    value = data.get("payload", data)
    return value if isinstance(value, dict) else {}


def _requirement_number(value: Any) -> float | None:
    if isinstance(value, (int, float)):
        return float(value)
    text_value = str(value or "").lower().replace(",", "")
    match = re.search(r"(\d+(?:\.\d+)?)", text_value)
    if not match:
        return None
    number = float(match.group(1))
    if "crore" in text_value or re.search(r"\bcr\b", text_value):
        number *= 10_000_000
    elif "lakh" in text_value or "lac" in text_value:
        number *= 100_000
    return number


_EVIDENCE_TERMS = {
    "general_experience": ["general experience", "construction experience"],
    "specific_experience_value": ["specific experience", "similar contract"],
    "specific_experience_count": ["specific experience", "number of contract"],
    "avg_annual_turnover": ["average annual construction turnover", "annual turnover"],
    "liquid_assets": ["liquid assets", "credit facilities"],
    "min_tender_capacity": ["minimum tender capacity", "tender capacity"],
    "tender_security": ["tender security"],
    "performance_security": ["performance security"],
    "retention_money": ["retention money", "retention"],
    "completion_time": ["completion time", "intended completion"],
    "equipment": ["equipment"],
    "personnel": ["personnel", "key personnel"],
    "payment_terms": ["payment terms", "payment certificate"],
}


def _pdf_requirement_evidence(path: str, requirements: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Locate each extracted criterion on a PDF page and return a short clause."""
    try:
        import fitz
        document = fitz.open(path)
    except Exception:
        return {}
    evidence: dict[str, dict[str, Any]] = {}
    try:
        pages = [(index + 1, page.get_text("text") or "") for index, page in enumerate(document)]
        for requirement in requirements:
            key = str(requirement.get("key") or "")
            raw_value = str(requirement.get("value") or "").strip()
            terms = []
            if len(raw_value) >= 5:
                terms.append(raw_value)
            terms.extend(_EVIDENCE_TERMS.get(key, [requirement.get("label", "")]))
            for page_number, page_text in pages:
                normalized = re.sub(r"\s+", " ", page_text)
                lower = normalized.lower()
                match_start = -1
                matched_text = ""
                for term in terms:
                    term_text = re.sub(r"\s+", " ", str(term or "")).strip()
                    if not term_text:
                        continue
                    match_start = lower.find(term_text.lower())
                    if match_start >= 0:
                        matched_text = normalized[match_start:match_start + len(term_text)]
                        break
                if match_start < 0:
                    continue
                start = max(0, match_start - 220)
                end = min(len(normalized), match_start + max(len(matched_text), 40) + 320)
                evidence[key] = {
                    "document_type": "TDS",
                    "page_number": page_number,
                    "clause_text": normalized[start:end].strip(),
                    "highlight_text": matched_text or str(terms[0]),
                    "match_method": "exact_text" if matched_text.lower() == raw_value.lower() else "criterion_keyword",
                }
                break
    finally:
        document.close()
    return evidence


def _text_requirement_evidence(path: str, requirements: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Locate criteria in page-preserving TDS text (`--- Page N ---` markers)."""
    try:
        text_value = Path(path).read_text(encoding="utf-8", errors="replace")
    except Exception:
        return {}
    parts = re.split(r"---\s*Page\s+(\d+)\s*---", text_value, flags=re.IGNORECASE)
    pages = [
        (int(parts[index]), parts[index + 1])
        for index in range(1, len(parts) - 1, 2)
        if parts[index].isdigit()
    ]
    evidence: dict[str, dict[str, Any]] = {}
    for requirement in requirements:
        key = str(requirement.get("key") or "")
        raw_value = str(requirement.get("value") or "").strip()
        terms = ([raw_value] if len(raw_value) >= 5 else []) + _EVIDENCE_TERMS.get(key, [requirement.get("label", "")])
        for page_number, page_text in pages:
            normalized = re.sub(r"\s+", " ", page_text).strip()
            lower = normalized.lower()
            for term in terms:
                term_text = re.sub(r"\s+", " ", str(term or "")).strip()
                match_start = lower.find(term_text.lower()) if term_text else -1
                if match_start < 0:
                    continue
                start = max(0, match_start - 220)
                end = min(len(normalized), match_start + len(term_text) + 320)
                evidence[key] = {
                    "document_type": "TDS",
                    "source_document_type": "tds_text",
                    "page_number": page_number,
                    "clause_text": normalized[start:end].strip(),
                    "highlight_text": normalized[match_start:match_start + len(term_text)],
                    "match_method": "page_preserved_text",
                }
                break
            if key in evidence:
                break
    return evidence


async def _state_entry(db: AsyncSession, tender_id: str) -> KnowledgeEntry | None:
    return (
        await db.execute(
            select(KnowledgeEntry)
            .where(
                KnowledgeEntry.tender_id == tender_id,
                KnowledgeEntry.entry_type == "tender_workspace_state",
                KnowledgeEntry.is_archived.is_(False),
            )
            .order_by(KnowledgeEntry.updated_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()


@router.get("/{tender_id}")
async def get_tender_workspace(
    tender_id: str,
    db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    """Return the complete 15-stage bid-preparation workspace."""
    lifecycle = (
        await db.execute(
            text("""
                SELECT package_no, title, agency_code, pe_office, procurement_method,
                       estimated_cost_bdt, award_amount_bdt, winner
                FROM procurement_lifecycle
                WHERE tender_id = :tid OR package_no = :tid
                ORDER BY (estimated_cost_bdt IS NOT NULL) DESC, updated_at DESC
                LIMIT 1
            """),
            {"tid": tender_id},
        )
    ).mappings().first()
    canonical = (
        await db.execute(
            text("""
                SELECT tender_id, package_no, title, agency_code, pe_office,
                       district, estimated_cost_bdt
                FROM canonical_tenders
                WHERE tender_id = :tid OR package_no = :tid
                ORDER BY rebuilt_at DESC LIMIT 1
            """),
            {"tid": tender_id},
        )
    ).mappings().first()
    live = (
        await db.execute(
            text("""
                SELECT tender_id, package_no, title, agency_code, closing_datetime
                FROM pf_tenders
                WHERE tender_id = :tid OR package_no = :tid
                ORDER BY closing_datetime DESC NULLS LAST LIMIT 1
            """),
            {"tid": tender_id},
        )
    ).mappings().first()

    knowledge_rows = (
        await db.execute(
            select(KnowledgeEntry)
            .where(KnowledgeEntry.tender_id == tender_id, KnowledgeEntry.is_archived.is_(False))
            .order_by(KnowledgeEntry.updated_at.desc())
        )
    ).scalars().all()
    by_type: dict[str, KnowledgeEntry] = {}
    for row in knowledge_rows:
        by_type.setdefault(row.entry_type, row)

    workspace_entry = await _state_entry(db, tender_id)
    state = dict(workspace_entry.data or {}) if workspace_entry else {}
    criteria = _payload(by_type.get("tds_criteria").data) if by_type.get("tds_criteria") else {}
    requirement_enrichment = (
        _payload(by_type.get("tender_requirement_enrichment").data)
        if by_type.get("tender_requirement_enrichment") else {}
    )
    document_knowledge = (
        _payload(by_type.get("tender_document").data) if by_type.get("tender_document") else {}
    )
    data_quality = (
        _payload(by_type.get("tender_data_quality").data) if by_type.get("tender_data_quality") else {}
    )
    document_paths = document_knowledge.get("documents") or {}

    # Include local tender-manager documents used by the upload workflow.
    from app.services.tender_manager import tender_manager

    local_documents: dict[str, str] = {}
    for doc_type in ("notice", "tds", "tds_2", "pcc", "gcc", "specifications", "drawings", "boq", "addendum", "clarification"):
        path = tender_manager.get_document_path(tender_id, doc_type)
        if path:
            local_documents[doc_type] = str(path)
    for key, value in document_paths.items():
        if value:
            local_documents.setdefault(str(key).lower().replace(" ", "_"), str(value))

    required_doc_names = [
        "Trade License", "TIN", "VAT", "Bank Solvency",
        "Similar Work Certificates", "Work Completion Certificates",
        "Experience Certificates", "Financial Statements", "Power of Attorney",
        "Tender Security", "Manufacturer Authorization", "JV Agreement",
    ]
    saved_statuses = state.get("document_statuses") or {}
    document_checklist = [
        {
            "name": name,
            "status": saved_statuses.get(name, "missing"),
            "required": name not in {"Manufacturer Authorization", "JV Agreement"},
        }
        for name in required_doc_names
    ]

    requirement_labels = {
        "general_experience": "General experience",
        "specific_experience_value": "Similar work value",
        "specific_experience_count": "Similar work count",
        "avg_annual_turnover": "Average annual turnover",
        "liquid_assets": "Liquid assets",
        "min_tender_capacity": "Minimum tender capacity",
        "tender_security": "Tender security",
        "performance_security": "Performance security",
        "retention_money": "Retention money",
        "completion_time": "Completion time",
        "equipment": "Equipment requirements",
        "personnel": "Personnel requirements",
        "payment_terms": "Payment terms",
    }
    requirements = [
        {"key": key, "label": requirement_labels.get(key, key.replace("_", " ").title()), "value": value,
         "source": "TDS", "status": "extracted"}
        for key, value in criteria.items()
        if value not in (None, "", [], {})
    ]
    tds_document = next(
        ((doc_type, path) for doc_type, path in local_documents.items() if "tds" in doc_type.lower() and Path(path).suffix.lower() == ".pdf"),
        None,
    )
    requirement_evidence = (
        await asyncio.to_thread(_pdf_requirement_evidence, tds_document[1], requirements)
        if tds_document else {}
    )
    tds_text_path = Path(__file__).resolve().parents[3] / "uploads" / tender_id / "tds_text.txt"
    if tds_text_path.is_file() and len(requirement_evidence) < len(requirements):
        text_evidence = await asyncio.to_thread(_text_requirement_evidence, str(tds_text_path), requirements)
        for key, value in text_evidence.items():
            requirement_evidence.setdefault(key, value)
    for requirement in requirements:
        evidence = requirement_evidence.get(requirement["key"], {
            "document_type": "TDS",
            "page_number": None,
            "clause_text": None,
            "highlight_text": str(requirement["value"]),
            "match_method": "structured_extraction",
        })
        evidence.update({
            "source_entry_type": "tds_criteria",
            "knowledge_entry_id": by_type.get("tds_criteria").id if by_type.get("tds_criteria") else None,
            "document_url": (
                f"/api/tender-workspace/{tender_id}/documents/tds_text"
                if evidence.get("source_document_type") == "tds_text"
                else f"/api/tender-workspace/{tender_id}/documents/{tds_document[0]}"
                if tds_document else None
            ),
        })
        requirement["evidence"] = evidence

    boq_stats = (
        await db.execute(
            text("""
                SELECT bc.id, bc.tender_id, bc.boq_file_id, bc.sor_agency, bc.zone,
                       bc.total_items, bc.matches, bc.variances, bc.mismatches,
                       bc.below_sor, bc.total_sor_amount, bc.total_quoted_amount,
                       bc.discount_pct, bc.excel_path, bc.docx_path, bc.updated_at
                FROM boq_comparisons bc
                LEFT JOIN tenders t ON t.id = bc.tender_id
                WHERE t.tender_id = :tid
                   OR bc.boq_file_id = :tid
                   OR bc.boq_file_id = :brain_tid
                ORDER BY bc.updated_at DESC LIMIT 1
            """),
            {"tid": tender_id, "brain_tid": f"brain-{tender_id}"},
        )
    ).mappings().first()
    boq_items: list[dict[str, Any]] = []
    if boq_stats and boq_stats["tender_id"]:
        rows = (
            await db.execute(
                text("""
                    SELECT item_no, code, description, unit, quantity, quoted_rate,
                           sor_rate, sor_code, diff, pct_diff, flag, work_type,
                           section, agency, attributes
                    FROM boq_items
                    WHERE tender_id = :tender_pk
                    ORDER BY created_at, item_no
                """),
                {"tender_pk": boq_stats["tender_id"]},
            )
        ).mappings().all()
        for row in rows:
            quantity = float(row["quantity"] or 0)
            completed_rate = row["quoted_rate"] if row["quoted_rate"] is not None else row["sor_rate"]
            completed_amount = quantity * float(completed_rate or 0)
            attributes = row["attributes"] if isinstance(row["attributes"], dict) else {}
            boq_items.append({
                "item_no": row["item_no"],
                "code": row["code"],
                "description": row["description"],
                "unit": row["unit"],
                "quantity": quantity,
                "quoted_rate": row["quoted_rate"],
                "sor_rate": row["sor_rate"],
                "completed_rate": completed_rate,
                "completed_amount": completed_amount,
                "sor_code": row["sor_code"],
                "agency": row["agency"],
                "match_type": attributes.get("match_type"),
                "match_confidence": attributes.get("match_confidence"),
                "flag": row["flag"],
                "unmatched_reason": (
                    attributes.get("remarks") or "No matching SOR code or description"
                    if row["sor_rate"] is None else None
                ),
            })
    qualification = (
        await db.execute(
            text("""
                SELECT qualification_score, recommendation, confidence_pct,
                       factors, risk_factors, explanation, created_at
                FROM tender_qualification_scores
                WHERE tender_id = :tid
                ORDER BY created_at DESC LIMIT 1
            """),
            {"tid": tender_id},
        )
    ).mappings().first()
    selected_contractor_id = state.get("contractor_id")
    if not selected_contractor_id:
        selected_contractor_id = await db.scalar(text("""
            SELECT contractor_id FROM contractor_documents
            GROUP BY contractor_id ORDER BY COUNT(*) DESC LIMIT 1
        """))
    contractor_docs = []
    contractor_profile: dict[str, Any] = {}
    if selected_contractor_id:
        contractor_docs = (await db.execute(text("""
            SELECT id, doc_category, doc_type, filename, extracted_data, created_at, updated_at
            FROM contractor_documents WHERE contractor_id=:contractor_id
            ORDER BY updated_at DESC NULLS LAST, created_at DESC
        """), {"contractor_id": selected_contractor_id})).mappings().all()
        profile_value = await db.scalar(text("""
            SELECT config FROM tenants
            WHERE config->>'contractor_db_id'=:contractor_id AND is_active=true
            ORDER BY (COALESCE(config->>'client_rank','999999'))::int LIMIT 1
        """), {"contractor_id": selected_contractor_id})
        contractor_profile = profile_value if isinstance(profile_value, dict) else {}

    document_terms = {
        "general_experience": ("general experience", "experience"),
        "specific_experience_value": ("specific experience", "similar work", "completion"),
        "specific_experience_count": ("specific experience", "similar work", "completion"),
        "avg_annual_turnover": ("turnover", "financial statement", "audit"),
        "liquid_assets": ("liquid asset", "credit commitment", "bank solvency"),
        "min_tender_capacity": ("tender capacity", "turnover", "work commitment"),
        "equipment": ("equipment", "machinery"),
        "personnel": ("personnel", "engineer", "cv", "resume"),
    }
    qualification_factors = qualification["factors"] if qualification and isinstance(qualification["factors"], dict) else {}
    eligibility_matrix = []
    for requirement in requirements:
        key = requirement["key"]
        company_value = qualification_factors.get(key)
        matched_docs = []
        for doc in contractor_docs:
            haystack = f"{doc['doc_category']} {doc['doc_type']} {doc['filename']}".lower()
            if any(term in haystack for term in document_terms.get(key, ())):
                years = [int(year) for year in re.findall(r"20\d{2}", doc["filename"] or "")]
                expired = bool(years and max(years) < datetime.now(timezone.utc).year - 1)
                matched_docs.append({
                    "document_id": doc["id"], "filename": doc["filename"],
                    "category": doc["doc_category"],
                    "updated_at": doc["updated_at"].isoformat() if doc["updated_at"] else None,
                    "expiry_status": "possibly_expired" if expired else "current_or_unstated",
                    "validity": (
                        doc["extracted_data"].get("validity")
                        if isinstance(doc["extracted_data"], dict) else None
                    ),
                })
        extracted_records = [
            doc["extracted_data"] for doc in contractor_docs
            if isinstance(doc["extracted_data"], dict) and doc["extracted_data"].get("verified")
        ]
        if company_value is None and key == "avg_annual_turnover":
            company_value = next((record.get("average_annual_turnover_bdt") for record in extracted_records if record.get("average_annual_turnover_bdt")), None)
        elif company_value is None and key == "specific_experience_value":
            values = [record.get("contract_value_bdt") for record in extracted_records if record.get("contract_value_bdt")]
            if contractor_profile.get("bwdb_similar_work_max_bdt"):
                values.append(contractor_profile["bwdb_similar_work_max_bdt"])
            company_value = max(values) if values else None
        elif company_value is None and key == "specific_experience_count":
            company_value = sum(1 for record in extracted_records if record.get("document_kind") == "similar_work_contract") or None
        elif company_value is None and key == "equipment":
            company_value = next((record.get("key_equipment") for record in extracted_records if record.get("document_kind") == "equipment_register"), None)
        elif company_value is None and key == "personnel":
            company_value = next((record.get("key_personnel_count") for record in extracted_records if record.get("key_personnel_count")), None)
        elif company_value is None and key == "liquid_assets":
            company_value = next((record.get("amount_bdt") for record in extracted_records if record.get("document_kind") == "credit_commitment"), None)
        elif company_value is None and key == "min_tender_capacity":
            company_value = next((record.get("assessed_tender_capacity_bdt") for record in extracted_records if record.get("document_kind") == "tender_capacity_assessment"), None)
        required_number = _requirement_number(requirement["value"])
        company_number = _requirement_number(company_value)
        invalid_evidence = any(
            doc.get("validity") in {"expired", "expired_contract_specific", "historical_recalculation_required"}
            for doc in matched_docs
        )
        status = (
            "pass" if company_value is True
            else "fail" if company_value is False
            else "expired_evidence" if invalid_evidence and company_value is not None
            else "pass" if required_number is not None and company_number is not None and company_number >= required_number
            else "fail" if required_number is not None and company_number is not None and company_number < required_number
            else "evidence_found" if matched_docs
            else "unknown"
        )
        eligibility_matrix.append({
            "criterion_key": key, "criterion": requirement["label"],
            "required_value": requirement["value"], "contractor_value": company_value,
            "status": status, "tender_evidence": requirement["evidence"],
            "company_evidence": matched_docs[:10],
            "company_profile_source": (
                contractor_profile.get("bwdb_similar_work_source")
                if key == "specific_experience_value" and contractor_profile.get("bwdb_similar_work_max_bdt")
                else None
            ),
            "expiry_status": (
                "expired_or_recalculation_required" if invalid_evidence
                else "possibly_expired" if any(doc["expiry_status"] == "possibly_expired" for doc in matched_docs)
                else "current_or_unstated" if matched_docs else "missing"
            ),
        })

    overview_source = lifecycle or canonical or live or {}
    from app.services.tender_estimate_service import TenderEstimateService
    estimate = await TenderEstimateService(db).resolve(
        tender_id=tender_id,
        package_no=overview_source.get("package_no") or document_knowledge.get("tender_info", {}).get("package_no"),
        tenant_id=user.get("tenant_id"),
    )
    if estimate["amount_bdt"] <= 0 and overview_source.get("estimated_cost_bdt"):
        estimate = {
            "amount_bdt": float(overview_source["estimated_cost_bdt"]),
            "source": "workspace_overview",
            "confidence": 0.85,
            "package_no": overview_source.get("package_no"),
        }
    closing = live.get("closing_datetime") if live else None
    now = datetime.now(timezone.utc)
    if closing and closing.tzinfo is None:
        closing = closing.replace(tzinfo=timezone.utc)
    countdown_seconds = max(0, int((closing - now).total_seconds())) if closing else None

    uploaded_required = sum(1 for item in document_checklist if item["status"] in {"uploaded", "verified"})
    required_total = sum(1 for item in document_checklist if item["required"])
    document_score = round(uploaded_required / max(required_total, 1) * 100)
    requirement_score = min(100, round(len(requirements) / 9 * 100))
    boq_completion = (
        round(sum(1 for item in boq_items if item["sor_rate"] is not None) / len(boq_items) * 100)
        if boq_items else int(state.get("boq_completion") or 0)
    )
    compliance_score = int(qualification["qualification_score"] * 100) if qualification else requirement_score
    readiness = 0

    tasks = state.get("tasks") or [
        {"id": "review", "title": "Review tender and TDS", "stage": "Review Tender", "status": "done" if requirements else "todo", "owner": state.get("owner") or "Bid Manager"},
        {"id": "eligibility", "title": "Verify eligibility", "stage": "Verify Eligibility", "status": "done" if qualification else "todo", "owner": "Compliance Lead"},
        {"id": "documents", "title": "Collect mandatory documents", "stage": "Collect Documents", "status": "done" if document_score == 100 else "in_progress", "owner": "Document Controller"},
        {"id": "forms", "title": "Prepare and verify forms", "stage": "Prepare Forms", "status": "todo", "owner": "Proposal Engineer"},
        {"id": "boq", "title": "Analyze and price BOQ", "stage": "Analyze & Price BOQ", "status": "done" if boq_completion == 100 else "todo", "owner": "Estimator"},
        {"id": "validation", "title": "Validate compliance and package", "stage": "Validate Compliance", "status": "todo", "owner": "Bid Manager"},
    ]
    missing_tasks = [task["title"] for task in tasks if task.get("status") != "done"]
    critical_blockers = [
        item["name"] for item in document_checklist
        if item["required"] and item["status"] == "missing"
    ]
    if not requirements:
        critical_blockers.insert(0, "TDS requirements not extracted")
    if not boq_stats:
        critical_blockers.append("BOQ analysis not completed")
    for issue in data_quality.get("issues") or []:
        if isinstance(issue, dict) and issue.get("severity") == "critical":
            critical_blockers.append(str(issue.get("message") or issue.get("code") or "Data quality failure"))

    report = _payload(by_type.get("executive_decision").data) if by_type.get("executive_decision") else {}
    agent_runs = (
        await db.execute(
            text("""
                SELECT id, agent_id, agent_name, status, error, output, execution_time_ms, created_at
                FROM agent_results
                WHERE tender_id = :tid AND agent_id NOT LIKE 'test-%'
                ORDER BY created_at DESC
                LIMIT 30
            """),
            {"tid": tender_id},
        )
    ).mappings().all()
    runtime = _payload(by_type.get("tender_pipeline_runtime").data) if by_type.get("tender_pipeline_runtime") else {}
    runtime_artifacts = runtime.get("artifacts") if isinstance(runtime.get("artifacts"), list) else []
    profit_data = _payload(by_type.get("profit_margin_analysis").data) if by_type.get("profit_margin_analysis") else {}
    profit_artifact = profit_data.get("artifact")
    if isinstance(profit_artifact, dict) and profit_artifact.get("path"):
        known_hashes = {item.get("sha256") for item in runtime_artifacts if isinstance(item, dict)}
        if profit_artifact.get("sha256") not in known_hashes:
            runtime_artifacts = [*runtime_artifacts, profit_artifact]
    forms = [
        "Letter of Tender", "Tender Submission Sheet", "Qualification Forms",
        "Experience Forms", "Personnel Forms", "Equipment Forms", "Price Schedule",
        "Manufacturer Forms", "Integrity Declaration",
    ]
    saved_form_statuses = state.get("form_statuses") or {}
    form_rows = [
        {
            "name": name,
            "status": saved_form_statuses.get(
                name,
                "ready" if name.lower().replace(" ", "_") in local_documents else "missing_fields",
            ),
        }
        for name in forms
    ]
    validation_checks = [
        ("Pre-agent data quality gate", data_quality.get("status") in {"passed", "warning"}),
        ("Missing documents", not critical_blockers),
        ("Missing signatures", False),
        ("Incorrect dates", closing is not None),
        ("BOQ formula and item validation", bool(boq_stats and not boq_stats["mismatches"])),
        ("Form consistency", False),
        ("File size limits", bool(local_documents)),
        ("PDF quality", bool(local_documents)),
        ("Naming convention", bool(local_documents)),
    ]
    completed_tasks = sum(1 for task in tasks if task.get("status") == "done")
    task_score = round(completed_tasks / max(len(tasks), 1) * 100)
    ready_forms = sum(1 for form in form_rows if form["status"] in {"ready", "verified", "completed"})
    form_score = round(ready_forms / max(len(form_rows), 1) * 100)
    passed_validations = sum(1 for _, passed in validation_checks if passed)
    validation_score = round(passed_validations / max(len(validation_checks), 1) * 100)
    approvals = state.get("approvals") or []
    required_approval_stages = ["eligibility", "documents", "boq", "pricing", "compliance", "final"]
    approved_stages = {
        str(item.get("stage") or "final")
        for item in approvals if isinstance(item, dict)
        and str(item.get("status", "")).lower() in {"approved", "complete", "completed"}
    }
    approval_score = round(len(approved_stages.intersection(required_approval_stages)) / len(required_approval_stages) * 100)
    completed_agent_runs = [
        row for row in agent_runs
        if str(row["status"] or "").lower() in {"success", "completed", "complete", "done"}
    ]
    agent_score = (
        round(len(completed_agent_runs) / len(agent_runs) * 100)
        if agent_runs else 0
    )
    readiness_components = {
        "documents": {"score": document_score, "weight": 25},
        "requirements": {"score": requirement_score, "weight": 10},
        "forms": {"score": form_score, "weight": 15},
        "boq": {"score": boq_completion, "weight": 20},
        "tasks": {"score": task_score, "weight": 10},
        "validation": {"score": validation_score, "weight": 10},
        "approval": {"score": approval_score, "weight": 5},
        "agents": {"score": agent_score, "weight": 5},
    }
    readiness = round(sum(
        component["score"] * component["weight"] / 100
        for component in readiness_components.values()
    ))
    readiness_snapshot = {
        "score": readiness,
        "components": readiness_components,
        "recalculated_at": now.isoformat(),
        "source_counts": {
            "documents_ready": uploaded_required,
            "documents_required": required_total,
            "requirements": len(requirements),
            "forms_ready": ready_forms,
            "forms_total": len(form_rows),
            "boq_items": len(boq_items),
            "boq_matched": sum(1 for item in boq_items if item["sor_rate"] is not None),
            "tasks_done": completed_tasks,
            "tasks_total": len(tasks),
            "validations_passed": passed_validations,
            "validations_total": len(validation_checks),
            "approvals": len(approvals),
            "successful_agents": len(completed_agent_runs),
            "agent_runs": len(agent_runs),
        },
    }
    previous_readiness = state.get("readiness") if isinstance(state.get("readiness"), dict) else {}
    previous_signature = previous_readiness.get("source_counts")
    if previous_signature != readiness_snapshot["source_counts"] or previous_readiness.get("score") != readiness:
        history = list(state.get("readiness_history") or [])[-49:]
        history.append({
            "score": readiness,
            "at": now.isoformat(),
            "reason": "workspace_sources_changed",
            "components": readiness_components,
        })
        state["readiness"] = readiness_snapshot
        state["readiness_history"] = history
        checksum = hashlib.sha256(f"{tender_id}:{state}".encode()).hexdigest()
        if workspace_entry is None:
            workspace_entry = KnowledgeEntry(
                id=str(uuid.uuid4()),
                entry_type="tender_workspace_state",
                tender_id=tender_id,
                tenant_id=user.get("tenant_id"),
                title=f"Tender workspace {tender_id}",
                summary=f"Submission readiness {readiness}/100",
                source="tender_workspace",
                data=state,
                checksum=checksum,
                tags={"workflow": "bid_preparation", "works_only": True},
            )
            db.add(workspace_entry)
        else:
            workspace_entry.data = state
            workspace_entry.summary = f"Submission readiness {readiness}/100"
            workspace_entry.checksum = checksum
        await db.commit()
        await db.refresh(workspace_entry)

    return {
        "tender_id": tender_id,
        "overview": {
            "title": overview_source.get("title") or document_knowledge.get("tender_info", {}).get("title") or f"Tender {tender_id}",
            "procuring_entity": overview_source.get("pe_office") or document_knowledge.get("tender_info", {}).get("procuring_entity"),
            "agency": overview_source.get("agency_code"),
            "package_no": overview_source.get("package_no"),
            "method": overview_source.get("procurement_method"),
            "estimated_cost_bdt": float(estimate["amount_bdt"]),
            "estimate_source": estimate["source"],
            "estimate_confidence": estimate["confidence"],
            "estimate_candidates": estimate.get("candidates", []),
            "estimate_status": estimate.get("status", "resolved" if estimate["amount_bdt"] > 0 else "unresolved"),
            "tender_security_amount_bdt": requirement_enrichment.get("tender_security_amount_bdt"),
            "tender_security_text": requirement_enrichment.get("tender_security_text") or criteria.get("tender_security"),
            "requirements_evidence": requirement_enrichment.get("evidence") or (
                {"document_type": "TDS", "source_entry_type": "tds_criteria"}
                if criteria else None
            ),
            "closing_at": closing.isoformat() if closing else document_knowledge.get("tender_info", {}).get("deadline"),
            "ai_summary": (
                by_type.get("executive_decision").summary
                if by_type.get("executive_decision") and by_type.get("executive_decision").summary
                else "Run the executive-decision pipeline after requirements and BOQ validation."
            ),
        },
        "progress": {
            "completion_pct": readiness,
            "readiness_score": readiness,
            "compliance_score": compliance_score,
            "missing_tasks": missing_tasks,
            "critical_blockers": critical_blockers,
            "countdown_seconds": countdown_seconds,
            "next_action": missing_tasks[0] if missing_tasks else "Internal approval and final submission",
        },
        "documents": [{"type": key, "name": value.split("\\")[-1].split("/")[-1], "path": value} for key, value in local_documents.items()],
        "requirements": requirements,
        "eligibility": {
            "score": float(qualification["qualification_score"] * 100) if qualification else 0,
            "recommendation": qualification["recommendation"] if qualification else "Not checked",
            "explanation": qualification["explanation"] if qualification else "Run the eligibility agent with the contractor profile.",
            "checks": [
                {"name": name, "status": "pass" if qualification and qualification["qualification_score"] >= 0.6 else "pending",
                 "explanation": qualification["explanation"] if qualification else "Awaiting contractor evidence"}
                for name in ["Similar work", "Financial capacity", "Average turnover", "Equipment", "Engineers", "Litigation history", "Blacklist status", "JV requirements"]
            ],
            "matrix": eligibility_matrix,
            "contractor_id": selected_contractor_id,
            "company_documents_considered": len(contractor_docs),
        },
        "document_checklist": document_checklist,
        "forms": form_rows,
        "boq": {
            **(dict(boq_stats) if boq_stats else {
                "total_items": 0, "matches": 0, "variances": 0, "mismatches": 0,
                "below_sor": 0, "total_sor_amount": 0, "total_quoted_amount": 0,
            }),
            "items": boq_items,
            "matched_items": sum(1 for item in boq_items if item["sor_rate"] is not None),
            "unmatched_items": sum(1 for item in boq_items if item["sor_rate"] is None),
            "completed_total_bdt": sum(item["completed_amount"] for item in boq_items),
            "export_url": (
                f"/api/boq/export/{boq_stats['boq_file_id']}?format=xlsx"
                if boq_stats and boq_stats.get("excel_path") else None
            ),
        },
        "rate_analysis": {
            "sor_rates_available": await db.scalar(text("SELECT COUNT(*) FROM sor_rates")),
            "previous_awards_available": await db.scalar(text("SELECT COUNT(*) FROM award_records_v2 WHERE amount_bdt > 0")),
            "status": "ready" if boq_stats else "awaiting_boq",
        },
        "pricing": {
            "total_bid_value": float(boq_stats["total_quoted_amount"] or 0) if boq_stats else 0,
            "direct_cost": float(boq_stats["total_sor_amount"] or 0) if boq_stats else 0,
            "discount_pct": float(boq_stats["discount_pct"] or 0) if boq_stats else 0,
            "win_probability": report.get("score"),
            "scenarios": ["Aggressive", "Balanced", "Safe"],
        },
        "compliance": {"score": compliance_score, "requirements_linked": len(requirements), "violations": critical_blockers},
        "validation": [{"name": name, "passed": passed} for name, passed in validation_checks],
        "data_quality": data_quality or {
            "status": "not_checked", "score": 0, "agent_use_allowed": False,
            "critical_count": 0, "warning_count": 0, "issues": [],
        },
        "risks": [
            {"category": "Data quality", "level": "high" if data_quality.get("status") == "blocked" else "medium" if data_quality.get("status") in {None, "warning"} else "low", "issue": f"{data_quality.get('critical_count', 0)} critical / {data_quality.get('warning_count', 0)} warnings", "mitigation": "Resolve guardrail findings and recheck before downstream agents"},
            {"category": "Documents", "level": "high" if critical_blockers else "low", "issue": f"{len(critical_blockers)} blockers", "mitigation": "Resolve required documents before approval"},
            {"category": "Pricing", "level": "medium" if not boq_stats else "low", "issue": "BOQ comparison pending" if not boq_stats else "BOQ compared", "mitigation": "Run SOR, market and previous-award comparison"},
            {"category": "Schedule", "level": "high" if countdown_seconds is not None and countdown_seconds < 86400 else "medium", "issue": "Submission deadline", "mitigation": "Freeze documents 24 hours before submission"},
        ],
        "collaboration": {
            "tasks": tasks,
            "comments": state.get("comments") or [],
            "approvals": state.get("approvals") or [],
            "activity": [{"event": "Workspace synchronized from PG17", "at": now.isoformat()}],
            "approval_stages": [
                {"stage": stage, "status": "approved" if stage in approved_stages else "pending"}
                for stage in required_approval_stages
            ],
        },
        "submission": {
            "readiness_score": readiness,
            "readiness_components": readiness_components,
            "recalculated_at": readiness_snapshot["recalculated_at"],
            "history": state.get("readiness_history") or [],
            "files": [{"name": item["name"], "status": item["status"]} for item in document_checklist],
            "sequence": ["Technical documents", "Qualification forms", "BOQ and price schedule", "Tender security", "Final signed package"],
            "package_available": bool(local_documents),
            "approval_complete": len(approved_stages.intersection(required_approval_stages)) == len(required_approval_stages),
            "submission_locked": len(approved_stages.intersection(required_approval_stages)) != len(required_approval_stages),
            "post_submission": ["Save e-GP receipt", "Record submission hash", "Schedule opening follow-up"],
        },
        "enterprise": {
            "countdown_seconds": countdown_seconds,
            "readiness_score": readiness,
            "compliance_score": compliance_score,
            "ai_confidence": report.get("confidence_level") or "pending",
            "critical_alerts": len(critical_blockers),
            "missing_documents": sum(1 for item in document_checklist if item["status"] == "missing"),
            "boq_completion": boq_completion,
            "validation_status": "passed" if all(passed for _, passed in validation_checks) else "action_required",
            "estimated_profit": 0,
            "win_probability": report.get("score"),
            "owner": state.get("owner") or "Bid Manager",
            "last_updated": (workspace_entry.updated_at if workspace_entry else now).isoformat(),
        },
        "agent_runs": [
            {
                "run_id": row["id"],
                "agent_id": row["agent_id"],
                "agent_name": row["agent_name"] or row["agent_id"],
                "status": row["status"],
                "error": row["error"],
                "output": row["output"] or {},
                "execution_time_ms": row["execution_time_ms"] or 0,
                "timestamp": row["created_at"].isoformat() if row["created_at"] else None,
            }
            for row in agent_runs
        ],
        "outputs": {
            "runtime_summary": by_type.get("tender_pipeline_runtime").summary if by_type.get("tender_pipeline_runtime") else None,
            "pipeline_result": runtime.get("pipeline_result") or {},
            "artifacts": [
                {
                    **{key: artifact.get(key) for key in ("name", "kind", "size_bytes", "sha256")},
                    "download_url": f"/api/tender-workspace/{tender_id}/artifacts/{index}",
                }
                for index, artifact in enumerate(runtime_artifacts)
                if isinstance(artifact, dict) and artifact.get("path")
            ],
        },
    }


@router.post("/{tender_id}/boq-complete")
async def complete_tender_boq_rates(
    tender_id: str,
    request: BOQCompletionRequest,
    db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
):
    """Extract BOQ, complete rates from all three SORs, persist and export."""
    from app.services.boq_compare_service import (
        BOQCompareError,
        resolve_owner_user_id,
        run_brain_compare_flow,
    )

    try:
        result, comparison_id = await run_brain_compare_flow(
            db,
            tender_id=tender_id,
            sor_agency=request.sor_agency,
            zone=request.zone.upper() if request.zone else None,
            user_id=await resolve_owner_user_id(db, user),
        )
    except BOQCompareError as exc:
        raise HTTPException(status_code=exc.status_code, detail=str(exc))
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"BOQ rate completion failed: {exc}")

    items = result.get("data", []) if isinstance(result.get("data"), list) else []
    matched = sum(1 for item in items if item.get("sor_rate") is not None)
    unmatched = len(items) - matched
    digest = hashlib.sha256(
        f"{tender_id}:{comparison_id}:{len(items)}:{matched}".encode("utf-8")
    ).hexdigest()
    db.add(KnowledgeEntry(
        id=str(uuid.uuid4()),
        entry_type="boq_rate_completion",
        tender_id=tender_id,
        tenant_id=user.get("tenant_id"),
        title=f"BOQ rate completion {tender_id}",
        summary=f"{matched}/{len(items)} items matched to BWDB/LGED/PWD SOR; {unmatched} unmatched",
        source="tender_workspace",
        procurement_type="Works",
        tags={"works_only": True, "boq": True, "sor_agencies": ["BWDB", "LGED", "PWD"]},
        data={
            "comparison_id": comparison_id,
            "total_items": len(items),
            "matched_items": matched,
            "unmatched_items": unmatched,
            "summary": result.get("summary", {}),
            "excel_path": result.get("excel_path"),
            "agencies_compared": result.get("agencies_compared", ["BWDB", "PWD", "LGED"]),
        },
        checksum=digest,
    ))
    await db.commit()
    return {
        "success": True,
        "comparison_id": comparison_id,
        "total_items": len(items),
        "matched_items": matched,
        "unmatched_items": unmatched,
        "summary": result.get("summary", {}),
        "export_url": (
            f"/api/boq/export/brain-{tender_id}?format=xlsx"
            if result.get("excel_path") else None
        ),
    }


@router.post("/{tender_id}/resolve-estimate")
async def resolve_tender_estimate(
    tender_id: str,
    allow_live: bool = Query(True),
    db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    """Resolve the estimate from PG17, optionally falling back to live e-GP APP."""
    from app.services.tender_estimate_service import TenderEstimateService
    resolution = await TenderEstimateService(db).resolve(
        tender_id=tender_id,
        package_no=None,
        allow_live=allow_live,
        tenant_id=user.get("tenant_id"),
    )
    await db.commit()
    return {"success": resolution["amount_bdt"] > 0, "tender_id": tender_id, **resolution}


@router.post("/{tender_id}/resolve-estimate/approve")
async def approve_tender_estimate(
    tender_id: str, request: EstimateApprovalRequest,
    db: AsyncSession = Depends(get_async_session), user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    payload = {**request.model_dump(), "status": "resolved", "source": "approved_app_recovery",
               "confidence": 1.0, "approved_at": datetime.now(timezone.utc).isoformat()}
    db.add(KnowledgeEntry(
        id=str(uuid.uuid4()), entry_type="tender_estimate_resolution", tender_id=tender_id,
        tenant_id=user.get("tenant_id"), title=f"Approved APP estimate {tender_id}",
        summary=f"Approved APP estimate BDT {request.amount_bdt:,.2f}", source="tender_workspace",
        data={"payload": payload}, checksum=hashlib.sha256(repr(payload).encode()).hexdigest(),
        tags={"works_only": True, "approved": True},
    ))
    await db.commit()
    return {"success": True, "tender_id": tender_id, **payload}


@router.post("/{tender_id}/eligibility-matrix/snapshot")
async def snapshot_eligibility_matrix(
    tender_id: str, db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    workspace = await get_tender_workspace(tender_id, db, user)
    matrix = workspace["eligibility"].get("matrix", [])
    payload = {"matrix": matrix, "captured_at": datetime.now(timezone.utc).isoformat()}
    db.add(KnowledgeEntry(
        id=str(uuid.uuid4()), entry_type="eligibility_evidence_matrix", tender_id=tender_id,
        tenant_id=user.get("tenant_id"), title=f"Eligibility evidence matrix {tender_id}",
        summary=f"{len(matrix)} eligibility criteria linked to tender evidence",
        source="tender_workspace", data={"payload": payload},
        checksum=hashlib.sha256(repr(payload).encode()).hexdigest(),
        tags={"works_only": True, "eligibility": True, "evidence": True},
    ))
    await db.commit()
    return {"success": True, **payload}


@router.get("/{tender_id}/boq/unmatched-suggestions")
async def unmatched_boq_suggestions(
    tender_id: str, limit: int = Query(5, ge=1, le=20),
    db: AsyncSession = Depends(get_async_session), user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    rows = (await db.execute(text("""
        SELECT bi.id, bi.item_no, bi.description, bi.unit
        FROM boq_items bi JOIN boq_comparisons bc ON bc.tender_id = bi.tender_id
        LEFT JOIN tenders t ON t.id = bc.tender_id
        WHERE (t.tender_id=:tid OR bc.boq_file_id=:brain) AND bi.sor_rate IS NULL
        ORDER BY bi.created_at LIMIT 100
    """), {"tid": tender_id, "brain": f"brain-{tender_id}"})).mappings().all()
    output = []
    for row in rows:
        words = [word for word in re.findall(r"[A-Za-z]{4,}", row["description"] or "")[:3]]
        pattern = "%" + "%".join(words) + "%" if words else "%"
        candidates = (await db.execute(text("""
            SELECT id, agency, code, description, unit, zone_a, zone_b, zone_c, zone_d
            FROM sor_rates WHERE description ILIKE :pattern
            ORDER BY CASE WHEN unit=:unit THEN 0 ELSE 1 END, agency, code LIMIT :limit
        """), {"pattern": pattern, "unit": row["unit"], "limit": limit})).mappings().all()
        output.append({**dict(row), "candidates": [dict(candidate) for candidate in candidates]})
    return {"success": True, "tender_id": tender_id, "items": output}


@router.post("/{tender_id}/boq/resolve")
async def resolve_unmatched_boq_item(
    tender_id: str, request: BOQResolutionRequest,
    db: AsyncSession = Depends(get_async_session), user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    rate = (await db.execute(text("""
        SELECT id, agency, code, description, unit, zone_a, zone_b, zone_c, zone_d
        FROM sor_rates WHERE agency=:agency AND code=:code LIMIT 1
    """), {"agency": request.agency, "code": request.sor_code})).mappings().first()
    if not rate:
        raise HTTPException(status_code=404, detail="SOR rate not found")
    amount = float(rate[f"zone_{request.zone.lower()}"] or 0)
    updated = (await db.execute(text("""
        UPDATE boq_items SET sor_rate=:rate, sor_code=:code, agency=:agency,
          attributes=COALESCE(attributes, '{}'::jsonb) || CAST(:attributes AS jsonb), updated_at=now()
        WHERE id=:id RETURNING id, description, unit
    """), {"rate": amount, "code": request.sor_code, "agency": request.agency,
           "id": request.item_id, "attributes": '{"match_type":"user_approved_sor","match_confidence":1.0}'})).mappings().first()
    if not updated:
        raise HTTPException(status_code=404, detail="BOQ item not found")
    payload = {"boq_item_id": request.item_id, "sor_code": request.sor_code,
               "agency": request.agency, "zone": request.zone.upper(), "rate_bdt": amount,
               "description": updated["description"], "approved_at": datetime.now(timezone.utc).isoformat()}
    db.add(KnowledgeEntry(
        id=str(uuid.uuid4()), entry_type="boq_sor_learning", tender_id=tender_id,
        tenant_id=user.get("tenant_id"), title=f"Approved SOR mapping {request.sor_code}",
        summary=f"Mapped BOQ item to {request.agency} {request.sor_code}",
        source="tender_workspace", data={"payload": payload},
        checksum=hashlib.sha256(repr(payload).encode()).hexdigest(),
        tags={"works_only": True, "boq": True, "approved": True},
    ))
    await db.commit()
    return {"success": True, **payload}


@router.post("/{tender_id}/watch/scan")
async def scan_tender_changes(
    tender_id: str, db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    current = (await db.execute(text("""
        SELECT package_no, closing_datetime, title, updated_at FROM pf_tenders
        WHERE tender_id=:tid OR package_no=:tid ORDER BY updated_at DESC NULLS LAST LIMIT 1
    """), {"tid": tender_id})).mappings().first()
    documents = (await db.execute(select(KnowledgeEntry).where(
        KnowledgeEntry.tender_id == tender_id, KnowledgeEntry.entry_type == "tender_document",
        KnowledgeEntry.is_archived.is_(False)).order_by(KnowledgeEntry.updated_at.desc()).limit(1))).scalar_one_or_none()
    doc_payload = _payload(documents.data) if documents else {}
    snapshot = {
        "package_no": current["package_no"] if current else None,
        "closing_at": current["closing_datetime"].isoformat() if current and current["closing_datetime"] else None,
        "source_updated_at": current["updated_at"].isoformat() if current and current["updated_at"] else None,
        "document_names": sorted((doc_payload.get("documents") or {}).keys()),
    }
    prior = (await db.execute(select(KnowledgeEntry).where(
        KnowledgeEntry.tender_id == tender_id, KnowledgeEntry.entry_type == "tender_change_watch",
        KnowledgeEntry.is_archived.is_(False)).order_by(KnowledgeEntry.updated_at.desc()).limit(1))).scalar_one_or_none()
    prior_snapshot = _payload(prior.data).get("snapshot", {}) if prior else {}
    changes = [{"field": key, "before": prior_snapshot.get(key), "after": value}
               for key, value in snapshot.items() if prior and prior_snapshot.get(key) != value]
    payload = {"snapshot": snapshot, "changes": changes, "status": "changed" if changes else "baseline" if not prior else "unchanged",
               "scanned_at": datetime.now(timezone.utc).isoformat()}
    db.add(KnowledgeEntry(
        id=str(uuid.uuid4()), entry_type="tender_change_watch", tender_id=tender_id,
        tenant_id=user.get("tenant_id"), title=f"Deadline and addendum watch {tender_id}",
        summary=f"{len(changes)} deadline/document changes detected", source="tender_workspace",
        data={"payload": payload}, checksum=hashlib.sha256(repr(payload).encode()).hexdigest(),
        tags={"works_only": True, "watcher": True, "changed": bool(changes)},
    ))
    await db.commit()
    return {"success": True, **payload}


@router.post("/{tender_id}/approvals")
async def set_stage_approval(
    tender_id: str, request: StageApprovalRequest,
    db: AsyncSession = Depends(get_async_session), user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    entry = await _state_entry(db, tender_id)
    data = dict(entry.data or {}) if entry else {}
    approvals = [item for item in data.get("approvals", []) if item.get("stage") != request.stage]
    if request.status == "approved":
        approvals.append({"id": str(uuid.uuid4()), "stage": request.stage, "status": "approved",
                          "role": user.get("role") or "Bid Team", "note": request.note,
                          "approved_at": datetime.now(timezone.utc).isoformat()})
    data["approvals"] = approvals
    checksum = hashlib.sha256(repr(data).encode()).hexdigest()
    if entry:
        entry.data, entry.checksum = data, checksum
    else:
        db.add(KnowledgeEntry(id=str(uuid.uuid4()), entry_type="tender_workspace_state",
            tender_id=tender_id, tenant_id=user.get("tenant_id"), title=f"Tender workspace {tender_id}",
            summary="Persistent tender preparation state", source="tender_workspace",
            data=data, checksum=checksum, tags={"works_only": True, "workflow": "bid_preparation"}))
    await db.commit()
    return {"success": True, "tender_id": tender_id, "approvals": approvals}


@router.post("/{tender_id}/profit-margin-workbook")
async def generate_profit_margin_workbook(
    tender_id: str,
    request: ProfitMarginRequest,
    db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    """Calculate pricing and persist an auditable Excel workbook as a runtime artifact."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    direct = request.direct_cost_bdt
    overhead = direct * request.overhead_pct / 100
    contingency = direct * request.contingency_pct / 100
    subtotal = direct + overhead + contingency
    profit = subtotal * request.target_profit_pct / 100
    before_tax = subtotal + profit
    tax_vat = before_tax * request.tax_vat_pct / 100
    bid = before_tax + tax_vat
    effective_margin = (profit / bid * 100) if bid else 0

    output_dir = Path(__file__).resolve().parents[3] / "runtime" / "outputs" / tender_id
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"Profit_Margin_{tender_id}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Profit Margin"
    ws.append(["Tender Profit Margin Calculator", tender_id])
    ws.append(["Component", "Input %", "Amount BDT"])
    rows = [
        ("Direct cost", None, direct),
        ("Overhead", request.overhead_pct, overhead),
        ("Contingency", request.contingency_pct, contingency),
        ("Cost subtotal", None, subtotal),
        ("Target profit", request.target_profit_pct, profit),
        ("Bid before tax/VAT", None, before_tax),
        ("Tax/VAT", request.tax_vat_pct, tax_vat),
        ("Recommended total bid", None, bid),
        ("Effective profit margin on bid", effective_margin, None),
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 22
    for row in range(3, 12):
        ws.cell(row, 2).number_format = '0.00"%"'
        ws.cell(row, 3).number_format = '#,##0.00'
    wb.save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    artifact = {
        "name": path.name,
        "kind": "profit_margin_xlsx",
        "path": str(path.resolve()),
        "size_bytes": path.stat().st_size,
        "sha256": digest,
    }

    runtime_entry = (
        await db.execute(
            select(KnowledgeEntry)
            .where(
                KnowledgeEntry.tender_id == tender_id,
                KnowledgeEntry.entry_type == "tender_pipeline_runtime",
                KnowledgeEntry.is_archived.is_(False),
            )
            .order_by(KnowledgeEntry.updated_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()
    runtime_data = _payload(runtime_entry.data) if runtime_entry else {}
    artifacts = [item for item in runtime_data.get("artifacts", []) if item.get("kind") != "profit_margin_xlsx"]
    artifacts.append(artifact)
    runtime_data["artifacts"] = artifacts
    if runtime_entry is None:
        runtime_entry = KnowledgeEntry(
            id=str(uuid.uuid4()), entry_type="tender_pipeline_runtime", tender_id=tender_id,
            tenant_id=user.get("tenant_id"), title=f"Tender runtime {tender_id}",
            summary="Persisted tender calculation output", source="profit_margin_calculator",
            data=runtime_data, checksum=digest, tags={"works_only": True, "runtime": True},
        )
        db.add(runtime_entry)
    else:
        runtime_entry.data = dict(runtime_data)
        runtime_entry.checksum = digest

    db.add(KnowledgeEntry(
        id=str(uuid.uuid4()), entry_type="profit_margin_analysis", tender_id=tender_id,
        tenant_id=user.get("tenant_id"), title=f"Profit margin analysis {tender_id}",
        summary=f"Recommended bid BDT {bid:,.2f}; effective margin {effective_margin:.2f}%",
        source="tender_workspace", data={
            **request.model_dump(), "overhead_bdt": overhead, "contingency_bdt": contingency,
            "profit_bdt": profit, "tax_vat_bdt": tax_vat, "recommended_bid_bdt": bid,
            "effective_margin_pct": effective_margin, "artifact": artifact,
        }, checksum=digest, tags={"works_only": True, "pricing": True},
    ))
    await db.commit()
    artifact_index = len(artifacts) - 1
    return {
        "success": True, "recommended_bid_bdt": bid, "profit_bdt": profit,
        "effective_margin_pct": effective_margin,
        "artifact": {**artifact, "download_url": f"/api/tender-workspace/{tender_id}/artifacts/{artifact_index}"},
    }


@router.get("/{tender_id}/artifacts/{artifact_index}")
async def download_tender_artifact(
    tender_id: str,
    artifact_index: int,
    db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
):
    """Download a persisted runtime artifact after containment validation."""
    entry = (
        await db.execute(
            select(KnowledgeEntry)
            .where(
                KnowledgeEntry.tender_id == tender_id,
                KnowledgeEntry.entry_type == "tender_pipeline_runtime",
                KnowledgeEntry.is_archived.is_(False),
            )
            .order_by(KnowledgeEntry.updated_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()
    payload = _payload(entry.data) if entry else {}
    artifacts = payload.get("artifacts") if isinstance(payload.get("artifacts"), list) else []
    profit_entry = (
        await db.execute(
            select(KnowledgeEntry)
            .where(
                KnowledgeEntry.tender_id == tender_id,
                KnowledgeEntry.entry_type == "profit_margin_analysis",
                KnowledgeEntry.is_archived.is_(False),
            )
            .order_by(KnowledgeEntry.updated_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()
    profit_artifact = _payload(profit_entry.data).get("artifact") if profit_entry else None
    if isinstance(profit_artifact, dict) and profit_artifact.get("path"):
        known_hashes = {item.get("sha256") for item in artifacts if isinstance(item, dict)}
        if profit_artifact.get("sha256") not in known_hashes:
            artifacts = [*artifacts, profit_artifact]
    if artifact_index < 0 or artifact_index >= len(artifacts):
        raise HTTPException(status_code=404, detail="Artifact not found")
    artifact = artifacts[artifact_index]
    path = Path(str(artifact.get("path") or "")).resolve()
    from app.core.config import settings
    backend_root = Path(__file__).resolve().parents[3]
    allowed_roots = [Path(settings.BASE_DIR).resolve(), (backend_root / "uploads").resolve()]
    if not path.is_file() or not any(root == path or root in path.parents for root in allowed_roots):
        raise HTTPException(status_code=404, detail="Artifact file is unavailable")
    return FileResponse(path, filename=str(artifact.get("name") or path.name))


@router.get("/{tender_id}/documents/{document_type}")
async def view_tender_evidence_document(
    tender_id: str,
    document_type: str,
    db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
):
    """Serve an acquired tender document for the authenticated evidence viewer."""
    allowed_types = {
        "notice", "tds", "tds_2", "pcc", "gcc", "specifications",
        "drawings", "boq", "addendum", "clarification", "tds_text",
    }
    normalized_type = document_type.lower().replace(" ", "_")
    if normalized_type not in allowed_types:
        raise HTTPException(status_code=404, detail="Document type not available")

    from app.services.tender_manager import tender_manager
    candidate = (
        Path(__file__).resolve().parents[3] / "uploads" / tender_id / "tds_text.txt"
        if normalized_type == "tds_text" else tender_manager.get_document_path(tender_id, normalized_type)
    )
    if not candidate:
        entry = (
            await db.execute(
                select(KnowledgeEntry)
                .where(
                    KnowledgeEntry.tender_id == tender_id,
                    KnowledgeEntry.entry_type == "tender_document",
                    KnowledgeEntry.is_archived.is_(False),
                )
                .order_by(KnowledgeEntry.updated_at.desc())
                .limit(1)
            )
        ).scalar_one_or_none()
        documents = _payload(entry.data).get("documents", {}) if entry else {}
        if isinstance(documents, dict):
            candidate = documents.get(normalized_type) or documents.get(document_type)
    if not candidate:
        raise HTTPException(status_code=404, detail="Acquired document not found")

    path = Path(str(candidate)).resolve()
    backend_root = Path(__file__).resolve().parents[3]
    permitted_roots = [
        (backend_root / "uploads").resolve(),
        (backend_root / "runtime" / "tender_acquisition").resolve(),
    ]
    if not path.is_file() or not any(path == root or root in path.parents for root in permitted_roots):
        raise HTTPException(status_code=404, detail="Document path is unavailable")
    return FileResponse(
        path,
        media_type="application/pdf" if path.suffix.lower() == ".pdf" else None,
        content_disposition_type="inline",
        filename=path.name,
    )


@router.patch("/{tender_id}")
async def update_tender_workspace(
    tender_id: str,
    update: WorkspaceStateUpdate,
    db: AsyncSession = Depends(get_async_session),
    user: dict = Depends(get_current_user),
) -> dict[str, Any]:
    """Persist preparation tasks, document statuses, approvals and comments."""
    entry = await _state_entry(db, tender_id)
    data = dict(entry.data or {}) if entry else {}
    for key, value in update.model_dump(exclude_none=True).items():
        data[key] = value
    checksum = hashlib.sha256(f"{tender_id}:{data}".encode()).hexdigest()
    if entry is None:
        entry = KnowledgeEntry(
            id=str(uuid.uuid4()),
            entry_type="tender_workspace_state",
            tender_id=tender_id,
            tenant_id=user.get("tenant_id"),
            title=f"Tender workspace {tender_id}",
            summary="Persistent tender preparation state",
            source="tender_workspace",
            data=data,
            checksum=checksum,
            tags={"workflow": "bid_preparation", "works_only": True},
        )
        db.add(entry)
    else:
        entry.data = data
        entry.checksum = checksum
    await db.commit()
    return {"success": True, "tender_id": tender_id, "state": data}
