"""Client-specific submission package generation for M/S. Hassan & Brothers.

The source templates are verified completed tender documents.  This module keeps
their layout while replacing tender identity, Notice/TDS requirements and BOQ
content with the current tender's data.
"""

from __future__ import annotations

import calendar
import shutil
import tempfile
import zipfile
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

from docx import Document

from app.services.boq_excel_generator import generate_boq_excel


TEMPLATE_DIR = Path(__file__).with_name("client_templates") / "hassan_and_brothers"
CLIENT_NAME = "M/S. Hassan & Brothers"
CLIENT_EMAIL = "hbl.engr@gmail.com"

DEFAULT_EQUIPMENT = [
    {"type": "Concrete Mixer Machine", "min_required": "4 Nos."},
    {"type": "Vibrator", "min_required": "4 Nos."},
    {"type": "Steel Formwork", "min_required": "3,000 sqm"},
    {"type": "Excavator", "min_required": "2 Nos."},
    {"type": "Engine Boat", "min_required": "4 Nos."},
    {"type": "Barge", "min_required": "2 Nos."},
]
DEFAULT_MANPOWER = [
    {"post": "Project Manager", "qualification": "BSc Civil Engineering", "nos": "1", "total_exp": "5 years", "similar_exp": "3 years"},
    {"post": "Site Engineer", "qualification": "BSc Civil Engineering", "nos": "1", "total_exp": "5 years", "similar_exp": "3 years"},
    {"post": "Site Engineer", "qualification": "Diploma in Civil Engineering", "nos": "2", "total_exp": "5 years", "similar_exp": "3 years"},
    {"post": "Surveyor", "qualification": "Diploma in Surveying/Civil", "nos": "1", "total_exp": "5 years", "similar_exp": "3 years"},
    {"post": "Foreman", "qualification": "Relevant trade experience", "nos": "2", "total_exp": "5 years", "similar_exp": "3 years"},
    {"post": "Safety Officer", "qualification": "Safety training and Works experience", "nos": "1", "total_exp": "3 years", "similar_exp": "2 years"},
]


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    """Replace paragraph text while retaining the first run's formatting."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _all_paragraphs(doc: Document) -> Iterable[Any]:
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        for container in (section.header, section.footer):
            yield from container.paragraphs
            for table in container.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from cell.paragraphs


def _replace_text_everywhere(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in _all_paragraphs(doc):
        text = paragraph.text
        updated = text
        for old, new in replacements.items():
            updated = updated.replace(old, new)
        if updated != text:
            _set_paragraph_text(paragraph, updated)


def _normalize_requirement(item: dict[str, Any], kind: str, index: int) -> dict[str, str]:
    if kind == "equipment":
        return {
            "sl": str(item.get("sl") or index),
            "type": str(item.get("type") or item.get("equipment") or item.get("name") or item.get("description") or ""),
            "min_required": str(item.get("min_required") or item.get("quantity") or item.get("qty") or item.get("nos") or ""),
        }
    return {
        "sl": str(item.get("sl") or index),
        "post": str(item.get("post") or item.get("position") or item.get("designation") or item.get("name") or ""),
        "qualification": str(item.get("qualification") or item.get("education") or ""),
        "nos": str(item.get("nos") or item.get("quantity") or item.get("qty") or "1"),
        "total_exp": str(item.get("total_exp") or item.get("experience") or item.get("total_experience") or ""),
        "similar_exp": str(item.get("similar_exp") or item.get("specific_experience") or ""),
    }


def _replace_table_rows(table: Any, rows: list[dict[str, str]], columns: list[str]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for values in rows:
        new_tr = deepcopy(table.rows[0]._tr)
        table._tbl.append(new_tr)
        cells = table.rows[-1].cells
        for index, key in enumerate(columns):
            cells[index].text = values.get(key, "")


def _common_replacements(data: Any) -> dict[str, str]:
    description = data.package_description or data.project_name or "the Works"
    return {
        "HB-WEL JV.": CLIENT_NAME,
        "HB-WEL JV": CLIENT_NAME,
        "info@handbl.com": CLIENT_EMAIL,
        "211693": data.tender_id or "",
        "147922": data.tender_id or "",
        "e-GP-49/ADP/Barawlia/PW/2017-18": data.package_no or "",
        "Construction of Permanent Slope Protection Work along the Right Bank of Padma River from Km 50.750 to Km 51.000 = 250.00m at Barawlia in Upazila-Louhajong, District-Munshiganj under Dhaka O&M Division-II, BWDB, Dhaka during the year 2017-2018.": description,
    }


def _generate_docx(template_name: str, output: Path, data: Any, kind: str) -> None:
    doc = Document(str(TEMPLATE_DIR / template_name))
    _replace_text_everywhere(doc, _common_replacements(data))
    description = data.package_description or data.project_name or "the Works"

    if kind == "equipment":
        rows = [_normalize_requirement(x, kind, i) for i, x in enumerate(data.equipment_list or DEFAULT_EQUIPMENT, 1)]
        _replace_table_rows(doc.tables[0], rows, ["sl", "type", "min_required"])
        for paragraph in doc.paragraphs:
            if "do hereby give assurance" in paragraph.text.lower():
                _set_paragraph_text(paragraph, f'We do hereby give assurance that all equipment required by the TDS will be mobilized in time for “{description}” (Tender ID {data.tender_id}, Package {data.package_no}).')
                break
    elif kind == "manpower":
        rows = [_normalize_requirement(x, kind, i) for i, x in enumerate(data.manpower_list or DEFAULT_MANPOWER, 1)]
        _replace_table_rows(doc.tables[0], rows, ["sl", "post", "qualification", "nos", "total_exp", "similar_exp"])
        for paragraph in doc.paragraphs:
            if "do hereby give assurance" in paragraph.text.lower():
                _set_paragraph_text(paragraph, f'We do hereby give assurance that the qualified personnel required by the TDS will be assigned in time for “{description}” (Tender ID {data.tender_id}, Package {data.package_no}).')
                break
    else:
        for paragraph in doc.paragraphs:
            if paragraph.text.strip().upper().startswith("NATURE OF WORKS"):
                _set_paragraph_text(paragraph, f"NATURE OF WORKS:\n{description}\nTender ID: {data.tender_id} | Package: {data.package_no}")
                break
        if doc.tables:
            rows = [_normalize_requirement(x, "equipment", i) for i, x in enumerate(data.equipment_list or DEFAULT_EQUIPMENT, 1)]
            _replace_table_rows(doc.tables[0], rows, ["sl", "type", "min_required"])

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))


def _parse_date(value: str) -> datetime | None:
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value[:19], fmt)
        except ValueError:
            continue
    return None


def _work_activities(data: Any) -> list[str]:
    descriptions = [str(item.get("work_type") or item.get("description") or "").strip() for item in data.boq_items]
    unique = []
    for value in descriptions:
        if value and value.lower() not in {x.lower() for x in unique}:
            unique.append(value[:90])
    description = f"{data.package_description} {data.project_name}".lower()
    if any(term in description for term in ("bank protection", "river", "geobag", "cc block", "slope protection")):
        base = [
            "Site Preparation & Mobilization",
            "Procurement of Manpower, Equipment and Materials",
            "Earthworks",
            "Dumping & Placing Geo-Textile Bags",
            "Manufacturing of C.C. Blocks",
            "Supply of Sand Filter, Geo-Textile Filter and Khoa Filter",
            "Dumping & Placing C.C. Blocks",
            "Associated C.C. Works",
            "Demobilization & Site Handover",
        ]
        return base
    if unique:
        base = [
            "Site Preparation & Mobilization",
            "Procurement of Manpower, Equipment and Materials",
            *unique[:5],
            "Testing, Quality Control and Rectification",
            "Demobilization & Site Handover",
        ]
    else:
        base = [
            "Site Preparation & Mobilization",
            "Procurement of Manpower, Equipment and Materials",
            "Survey, Setting Out and Temporary Works",
            "Earthwork and Foundation Preparation",
            "Structural / Reinforced Concrete Works",
            "Associated Civil and Protection Works",
            "Finishing and Site Restoration",
            "Testing, Quality Control and Rectification",
            "Demobilization & Site Handover",
        ]
    return base[:9] + [""] * max(0, 9 - len(base))


def _patch_work_plan(data: Any, output: Path) -> None:
    """Patch cells through OOXML so the source workbook's shapes stay intact."""
    source = TEMPLATE_DIR / "Work_Plan_Template.xlsx"
    start = _parse_date(data.project_start_date) or ((_parse_date(data.tender_close_datetime) or datetime.now()) + timedelta(days=30))
    description = data.package_description or data.project_name or "the Works"
    values = {
        "B3": f"Tender ID: {data.tender_id}",
        "D3": "WORK SCHEDULE",
        "K3": f"Package: {data.package_no}",
        "B4": f"Name of Work: {description}",
    }
    for offset, ref in enumerate(["D9", "E9", "F9", "G9", "H9", "I9", "J9", "K9", "L9", "M9", "N9", "O9"]):
        month = ((start.month - 1 + offset) % 12) + 1
        year = start.year + (start.month - 1 + offset) // 12
        values[ref] = f"{calendar.month_abbr[month]} {year}"
    for row, activity in enumerate(_work_activities(data), 10):
        values[f"C{row}"] = activity

    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with tempfile.TemporaryDirectory() as td:
        expanded = Path(td) / "xlsx"
        with zipfile.ZipFile(source) as archive:
            archive.extractall(expanded)
        sheet_path = expanded / "xl" / "worksheets" / "sheet1.xml"
        sheet = ET.parse(sheet_path)
        for cell in sheet.getroot().findall(".//m:c", ns):
            ref = cell.attrib.get("r")
            if ref not in values:
                continue
            value_node = cell.find("m:v", ns)
            if value_node is None:
                continue
            # Inline strings are cell-local. Editing a shared-string entry can
            # unintentionally alter another cell that points at the same index.
            cell.attrib["t"] = "inlineStr"
            for child in list(cell):
                if child.tag in {f"{{{ns['m']}}}v", f"{{{ns['m']}}}is"}:
                    cell.remove(child)
            inline = ET.SubElement(cell, f"{{{ns['m']}}}is")
            text_node = ET.SubElement(inline, f"{{{ns['m']}}}t")
            text_node.text = values[ref]
        sheet.write(sheet_path, encoding="utf-8", xml_declaration=True)
        workbook_path = expanded / "xl" / "workbook.xml"
        workbook = ET.parse(workbook_path)
        sheet_node = workbook.getroot().find(".//m:sheet", ns)
        if sheet_node is not None:
            sheet_node.attrib["name"] = (data.tender_id or "Work Plan")[:31]
        workbook.write(workbook_path, encoding="utf-8", xml_declaration=True)
        output.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for path in expanded.rglob("*"):
                if path.is_file():
                    archive.write(path, path.relative_to(expanded).as_posix())


def generate_client_package(data: Any, output_dir: str) -> dict[str, str]:
    """Generate real client documents from Notice/TDS/BOQ data."""
    target = Path(output_dir)
    outputs: dict[str, str] = {}
    jobs = [
        ("4_Equipment_Declaration.docx", "Equipment_Declaration_Template.docx", "equipment"),
        ("5_Manpower_Declaration.docx", "Manpower_Declaration_Template.docx", "manpower"),
        ("6_Methodology.docx", "Methodology_Template.docx", "methodology"),
    ]
    for filename, template, kind in jobs:
        _generate_docx(template, target / filename, data, kind)
        outputs[filename] = "OK"
    _patch_work_plan(data, target / "7_Work_Plan.xlsx")
    outputs["7_Work_Plan.xlsx"] = "OK"

    tender_info = {
        "tender_id": data.tender_id,
        "package_no": data.package_no,
        "description": data.package_description or data.project_name,
        "procuring_entity": data.procuring_entity,
        "organization": data.organization,
        "estimated_value": data.estimated_value_bdt,
        "tender_security": data.tender_security_bdt,
        "closing_date": data.tender_close_datetime,
        "completion_period_days": data.completion_period_days,
    }
    # Always create the five-sheet analysis workbook. When the BOQ has not
    # extracted yet it is an honest empty analysis shell, never sample values.
    generate_boq_excel(tender_info, data.boq_items or [], output_path=str(target / "8_BOQ_Rate_Analysis.xlsx"))
    outputs["8_BOQ_Rate_Analysis.xlsx"] = "OK" if data.boq_items else "OK: awaiting BOQ extraction"
    return outputs
